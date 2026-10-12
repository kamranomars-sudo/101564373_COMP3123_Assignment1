from pathlib import Path
from io import BytesIO
from pypdf import PdfReader
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from PIL import Image

root = Path(__file__).resolve().parents[1]
pdf = PdfReader(root / 'evidence/comp3123_assignment1_screenshots.pdf')
titles = ['Postman collection test results', 'Postman main collection and rate limiting', 'Rate limiting and MongoDB users', 'MongoDB employee ownership for User A', 'MongoDB employee ownership for User B', 'Cross user retrieval and update rejection', 'Cross user update rejection', 'Cross user deletion rejection', 'Employee lists filtered by owner', 'Successful employee updates', 'Successful employee deletions', 'Input validation and missing authentication', 'Health check and Helmet security headers', 'Local application logs']
image_titles = ['Postman collection test results'] * 3 + ['Repeated signup conflict results', 'Authentication rate limiting', 'MongoDB users and hashed passwords', 'MongoDB employee ownership for User A', 'MongoDB employee ownership for User B'] + ['Cross user retrieval rejection'] * 2 + ['Cross user update rejection'] * 2 + ['Cross user deletion rejection'] * 2 + ['Employee lists filtered by owner'] * 2 + ['Successful employee updates'] * 2 + ['Successful employee deletions'] * 2 + ['Employee input validation', 'Missing authentication token', 'Health check', 'Helmet security headers', 'Local application logs']
doc = Document()
s = doc.sections[0]
s.orientation = WD_ORIENT.LANDSCAPE
s.page_width, s.page_height = Inches(11), Inches(8.5)
s.top_margin = s.bottom_margin = Inches(.5)
s.left_margin = s.right_margin = Inches(.55)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Arial'
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Heading 1'].font.size = Pt(16)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
doc.add_paragraph('COMP3123 Assignment 1 API Test Evidence', 'Title')
doc.add_paragraph('Student number 101564373')
doc.add_paragraph('Postman outputs and error responses, MongoDB records and ownership relationships, HTTP security headers, health checks and local application logs are recorded in the following screenshots.')
doc.add_paragraph('GitHub repository https://github.com/kamranomars-sudo/101564373_COMP3123_Assignment1')
count = 0
for page_index, page in enumerate(pdf.pages):
    for image_index, image in enumerate(page.images):
        count += 1
        doc.add_page_break()
        doc.add_paragraph(image_titles[count - 1], 'Heading 1')
        doc.add_paragraph(f'Screenshot {count}   Source evidence page {page_index + 1}')
        raw = image.data
        im = Image.open(BytesIO(raw))
        width = min(9.9, 6.05 * im.width / im.height)
        doc.add_picture(BytesIO(raw), width=Inches(width))
doc.core_properties.title = 'COMP3123 Assignment 1 API Test Evidence'
doc.core_properties.author = '101564373'
out = root / 'evidence/101564373_Postman_Evidence.docx'
doc.save(out)
print(f'Created DOCX with {count} screenshots')
